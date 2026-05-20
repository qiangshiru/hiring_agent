from pathlib import Path
from typing import BinaryIO

from src.services.resume_parser.loaders.base import (
    BaseDocumentLoader,
    CorruptedFileError,
    RawDocument,
)


class DOCXLoader(BaseDocumentLoader):
    supported_extensions = [".docx"]

    def load(self, file_path: str | Path) -> RawDocument:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        if not path.is_file():
            raise CorruptedFileError(f"Not a file: {file_path}")

        try:
            from docx import Document

            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n".join(paragraphs)

            tables_content: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text for cell in row.cells if cell.text.strip()]
                    if cells:
                        tables_content.append(" | ".join(cells))

            if tables_content:
                content += "\n\n[Tables]\n"
                content += "\n".join(tables_content)

            if not content.strip():
                raise CorruptedFileError(f"DOCX appears to be empty: {file_path}")

            return RawDocument(
                content=content,
                file_name=path.name,
                file_type="docx",
                metadata={
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables),
                },
            )
        except ImportError:
            raise CorruptedFileError("python-docx is required to load DOCX files")
        except Exception as exc:
            raise CorruptedFileError(f"Failed to load DOCX: {exc}")

    def load_from_stream(self, stream: BinaryIO, file_name: str) -> RawDocument:
        try:
            from docx import Document

            doc = Document(stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n".join(paragraphs)

            tables_content: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text for cell in row.cells if cell.text.strip()]
                    if cells:
                        tables_content.append(" | ".join(cells))

            if tables_content:
                content += "\n\n[Tables]\n"
                content += "\n".join(tables_content)

            if not content.strip():
                raise CorruptedFileError("DOCX appears to be empty")

            return RawDocument(
                content=content,
                file_name=file_name,
                file_type="docx",
                metadata={
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables),
                },
            )
        except ImportError:
            raise CorruptedFileError("python-docx is required to load DOCX files")
        except Exception as exc:
            raise CorruptedFileError(f"Failed to load DOCX from stream: {exc}")
